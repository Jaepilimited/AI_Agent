"""공유드라이브에서 롯트로 COA·MSDS 를 찾는다.

⛔ 이 모듈은 LLM 을 부르지 않는다. 파일 선택은 규칙이 한다 — 성분 조회를 LLM 에
   맡기지 않은 것과 같은 이유다. 규제 문서에서 틀린 파일을 주는 것은
   못 찾는 것보다 나쁘다.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Iterable, Optional, Sequence

# 헤더 낱말. ⛔ "몇 번째 행" 으로 박지 마라 — 안내문이 늘면 조용히 0건이 난다
_HEADER_ALIASES = {
    "sku": ("sku", "품목", "품목코드", "제품코드"),
    "description": ("description", "제품명", "품목명", "품명", "desc"),
    "lot": ("lot", "롯트", "로트", "lot no", "lot.no", "제조번호"),
}


@dataclass(frozen=True)
class Row:
    sku: str
    description: str
    lot: str
    line_no: int


@dataclass(frozen=True)
class ParsedList:
    """읽어들인 행과, **말없이 버린 행의 수**, 그리고 **어떻게 읽었는지**.

    ⛔ 버릴 거면 버렸다고 말해야 한다. 40행을 넣고 38행을 받아도 사용자는
       세어보지 않으면 모른다 — 에러가 아니라 조용한 유실이다.
    ⛔ 헤더 없이 알아본 경우도 마찬가지다. 어느 칸을 롯트로 봤는지 말하지 않으면
       틀리게 읽어도 아무도 모른다 (`layout`).
    """
    rows: tuple[Row, ...]
    skipped_no_sku: int = 0
    inferred: bool = False
    layout: str = ""


# 롯트 생김새 (실측 표본: F31C28 D · E08Z011 · FE103C · FE161 · MO388 · 2UE0003 · 6752FE)
_LOT_SHAPES = (
    re.compile(r"^[A-Za-z]\d{2}[A-Za-z]\d{2,3}$"),   # F31C28 · E08Z011 · C24Z068
    re.compile(r"^[A-Za-z]{2}\d{3}[A-Za-z]?$"),      # FE103C · FE161 · MO388
    re.compile(r"^\d[A-Za-z]{2}\d{4}$"),             # 2UE0003
    re.compile(r"^\d{4}[A-Za-z]{2}$"),               # 6752FE
)
# ⚠️ 순수 숫자는 수량·단가일 수도 있다 — 그 행에 더 나은 후보가 없을 때만 쓴다
_WEAK_LOT_SHAPE = re.compile(r"^\d{6}$")             # 416022 · 416006
_SKU_SHAPE = re.compile(r"^[A-Za-z]{3,6}\d{2,4}$")   # EUSKA022 · KRSKS003
# 2칸 이상 띄어쓰기만 칸을 가른다. ⛔ 홑 공백으로 쪼개면 제품명이 토막 난다
# ("Poremizing Fresh Ampoule 50ml")
_WIDE_GAP = re.compile(r" {2,}|\t")


def _is_lot(word: str) -> bool:
    return any(p.match(word) for p in _LOT_SHAPES)


def _find_lot(words: Sequence[str]) -> tuple[str, Optional[tuple[int, int]]]:
    """낱말들 속의 롯트와 그것이 차지한 구간. 접미('F31C28 D')를 함께 집는다."""
    for i, w in enumerate(words):
        if _is_lot(w):
            nxt = words[i + 1] if i + 1 < len(words) else ""
            if len(nxt) == 1 and nxt.isalpha():
                return f"{w} {nxt}", (i, i + 1)
            return w, (i, i)
    for i, w in enumerate(words):
        if _WEAK_LOT_SHAPE.match(w):
            return w, (i, i)
    return "", None


def _split_table(text: str) -> tuple[list[list[str]], str]:
    """구분자를 알아내 표로 쪼갠다. 엑셀 밖에서 복사한 것도 받는다."""
    lines = (text or "").splitlines()
    if any("\t" in ln for ln in lines):
        return [ln.split("\t") for ln in lines], "탭"
    filled = [ln for ln in lines if ln.strip()]
    if filled and all("," in ln for ln in filled):
        # "일관되게 쪼개질" 때만 쉼표로 본다 — 제품명에 쉼표가 있으면 칸 수가 흔들린다
        if len({ln.count(",") for ln in filled}) == 1:
            return [ln.split(",") for ln in lines], "쉼표"
    return [_WIDE_GAP.split(ln) for ln in lines], "띄어쓰기"


def _infer_row(cells: Sequence[str], line_no: int) -> Optional[Row]:
    """헤더가 없을 때 생김새로 열을 알아본다.

    ⛔ 알아본 결과를 조용히 쓰지 마라 — 무엇을 롯트로 봤는지 화면이 말한다
       (`ParsedList.layout`).
    """
    cells = [c.strip() for c in cells if c and c.strip()]
    if not cells:
        return None

    lot = sku = ""
    rest: list[str] = []
    for c in cells:
        words = c.split()
        found, span = _find_lot(words)
        whole = span is not None and span[0] == 0 and span[1] == len(words) - 1
        if not lot and whole:
            lot = found
            continue
        if not sku and _SKU_SHAPE.match(c):
            sku = c
            continue
        rest.append(c)

    # 칸이 홑 공백으로 붙어 온 경우에만 칸 **안**을 들여다본다
    if not lot:
        for i, c in enumerate(rest):
            words = c.split()
            found, span = _find_lot(words)
            if found and span is not None:
                lot = found
                rest[i] = " ".join(words[:span[0]] + words[span[1] + 1:])
                break
    if not sku:
        for i, c in enumerate(rest):
            words = c.split()
            for j, w in enumerate(words):
                if _SKU_SHAPE.match(w):
                    sku = w
                    rest[i] = " ".join(words[:j] + words[j + 1:])
                    break
            if sku:
                break

    description = max((r for r in rest if r), key=len, default="")
    if not (lot or sku or description):
        return None
    return Row(sku=sku, description=description, lot=lot, line_no=line_no)


def _match_header(cells: Sequence[str]) -> dict:
    """이 행에서 알아본 {필드: 열번호}. 아무것도 못 알아보면 빈 dict."""
    idx: dict[str, int] = {}
    for col, raw in enumerate(cells):
        key = (raw or "").strip().casefold().replace("_", " ")
        for field, aliases in _HEADER_ALIASES.items():
            if field not in idx and key in aliases:
                idx[field] = col
    return idx


def _looks_like_header(cells: Sequence[str]) -> Optional[dict]:
    """머리말 낱말로만 이뤄진 행이라야 헤더다.

    ⚠️ 알아보지 못한 칸이 하나라도 있으면 데이터 행으로 본다 — 안 그러면
       'LOT' 이라는 말이 섞인 데이터 행을 헤더로 삼는다.
    """
    hit = _match_header(cells)
    if not hit:
        return None
    filled = [c for c in cells if c and c.strip()]
    return hit if len(hit) >= len(filled) else None


def _cell(cells: Sequence[str], col: Optional[int]) -> str:
    if col is None or col >= len(cells):
        return ""
    return (cells[col] or "").strip()


def _rows_from_cells(table: Iterable[Sequence[str]],
                     sep_label: str = "") -> ParsedList:
    """헤더 행을 찾고 그 아래를 읽는다. 헤더가 없으면 생김새로 알아본다.

    ⛔ 헤더를 몇 번째 행이라고 박지 마라. 머리말 안내문이 한 줄만 늘어도 어긋나고,
       그때 나는 것은 에러가 아니라 **0건**이다 (OP 재고 적재에서 겪은 함정).
    ⛔ 알아볼 수 있는 것을 거절하지 마라 — 헤더 없이 붙여넣는 것도, 롯트만
       붙여넣는 것도 사람이 할 만한 일이다 (프로덕션 400 두 건의 원인).
    """
    table = [list(cells) for cells in table]
    header: Optional[dict] = None
    header_at = -1
    for i, cells in enumerate(table):
        hit = _looks_like_header(cells)
        if hit is not None:
            header, header_at = hit, i
            break

    rows: list[Row] = []
    skipped_no_sku = 0

    if header is None:
        # 헤더가 없다 — 생김새로 알아본다
        for line_no, cells in enumerate(table, start=1):
            row = _infer_row(cells, line_no)
            if row is not None:
                rows.append(row)
        layout = ""
        if rows:
            layout = (
                f"헤더 행이 없어 열을 알아봤습니다({sep_label or '자동'} 구분) — "
                f"롯트 {sum(1 for r in rows if r.lot)}건 · "
                f"SKU {sum(1 for r in rows if r.sku)}건 · "
                f"제품명 {sum(1 for r in rows if r.description)}건. "
                "다르게 읽혔으면 SKU·제품명·롯트 머리글을 붙여 다시 넣어주세요")
        return ParsedList(rows=tuple(rows), inferred=True, layout=layout)

    has_sku = "sku" in header
    for offset, cells in enumerate(table[header_at + 1:], start=1):
        sku = _cell(cells, header.get("sku"))
        lot = _cell(cells, header.get("lot"))
        description = _cell(cells, header.get("description"))
        if has_sku and not sku:
            # ⛔ 빈 줄을 거르는 가드가 "SKU 만 빈 행" 까지 삼켰다. 내용이 있는
            #    행을 버릴 때는 **세어서 알린다** — 빈 줄은 원래대로 조용히 버린다
            #    (그것까지 세면 숫자가 소음이 되어 아무도 안 본다)
            if lot or description:
                skipped_no_sku += 1
            continue
        if not (sku or lot or description):
            continue
        rows.append(Row(sku=sku, description=description, lot=lot,
                        line_no=header_at + 1 + offset))

    layout = ""
    if "lot" not in header:
        # ⛔ 조용히 진행하지 마라 — 롯트 열이 없으면 COA 롯트 매칭 자체가 불가능하다
        layout = ("붙여넣은 표에 롯트 열이 없습니다 — COA 는 롯트로 찾으므로 "
                  "이 목록은 제품명 기준으로만 조회합니다")
    return ParsedList(rows=tuple(rows), skipped_no_sku=skipped_no_sku, layout=layout)


def parse_pasted(text: str) -> ParsedList:
    """엑셀에서 복사한 표. 탭·쉼표·여러 칸 띄어쓰기를 모두 받는다."""
    table, sep_label = _split_table(text)
    return _rows_from_cells(table, sep_label=sep_label)


def parse_xlsx(data: bytes) -> ParsedList:
    """.xlsx 업로드. 첫 시트만 읽는다."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    table = [
        ["" if c is None else str(c) for c in row]
        for row in ws.iter_rows(values_only=True)
    ]
    return _rows_from_cells(table)


FOUND = "찾음"
MANY = "여러건"
NONE = "없음"
CHECK = "확인필요"
# ⛔ 조회 실패는 판정이 아니다. 확인필요로 뭉개면 "애매한 매칭 8건" 과
#    "조회가 아예 안 된 8건" 이 화면에서 똑같이 보인다 — 게다가 확인필요는
#    이미 다른 뜻(ZIP 에 경고를 달고 나가는 항목)을 지고 있다
FAILED = "조회실패"

_QUERY_FAILED_NOTE = "드라이브 조회에 실패했습니다 — 이 행은 판정하지 못했습니다."


def _failure_cause(exc: BaseException) -> str:
    """무엇을 하면 되는지까지 말한다. 한 문장 고정이면 429 와 권한 문제가
    사용자에게 똑같이 보인다.

    ⛔ 예외 원문을 붙이지 마라 — URL·토큰이 섞여 화면으로 샌다. 판정은
       HTTP 상태 코드로만 하고, 모르면 **모른다고** 한다 (아는 척이 더 나쁘다).
    """
    status = getattr(getattr(exc, "resp", None), "status", None)
    try:
        status = int(status)
    except (TypeError, ValueError):
        status = None
    if status in (401, 403):
        return " 접근 권한 문제로 보입니다 — 구글 연결과 이 드라이브의 멤버 여부를 확인하세요"
    if status == 429:
        return " 요청이 몰려 거절됐습니다 — 잠시 뒤 다시 시도하세요"
    if status is not None and 500 <= status < 600:
        return " 드라이브 쪽 일시 오류입니다 — 다시 시도하세요"
    return " 원인은 특정하지 못했습니다 — 다시 시도해도 같으면 관리자에게 알려주세요"

# 짧은 롯트는 다른 코드의 머리에 걸릴 수 있다 (Drive 는 토큰 앞부분 매칭이다)
_SHORT_LOT_LEN = 4

# '없음' 이 무엇을 근거로 한 없음인지. 빈 note 는 근거 없는 단정과 같다
_NONE_NOTE = ("파일명에 이 롯트가 든 파일이 없습니다 "
              "(파일 본문은 검색하지 않습니다)")

# 사본 표기. 원본과 같은 파일인데 이름만 다르다
_COPY_SUFFIX = re.compile(r"(의\s*사본|\s*\(\d+\)|\s*-\s*복사본)\s*$")


@dataclass(frozen=True)
class DriveFile:
    id: str
    name: str
    size: int
    web_link: str


@dataclass(frozen=True)
class Verdict:
    status: str
    files: tuple[DriveFile, ...]
    note: str = ""


def _dedup_key(f: DriveFile) -> tuple:
    base = _COPY_SUFFIX.sub("", f.name).strip()
    return (base.casefold(), f.size)


def _dedup(files: Sequence[DriveFile]) -> tuple[DriveFile, ...]:
    """같은 파일의 사본을 접는다. 먼저 온 것을 남긴다."""
    seen: dict[tuple, DriveFile] = {}
    for f in files:
        seen.setdefault(_dedup_key(f), f)
    return tuple(seen.values())


def _is_delimited_match(name: str, lot: str) -> bool:
    """lot 앞뒤가 영숫자와 붙어 있지 않은 매칭이 하나라도 있는가.

    ⛔ 경계를 안 보면 FE161 이 FE1615 에도 걸린다 — 다른 롯트의 증명서를
       주는 것은 못 찾는 것보다 나쁘다.
    ⛔ 한쪽만 casefold 하지 마라 — 자릿수 계산이 **검색한 그 문자열** 위에서
       돌아야 경계 판정이 맞는다. 둘 다 접어서 같은 문자열을 본다
    """
    name = name.casefold()
    lot = lot.casefold()
    start = 0
    while True:
        idx = name.find(lot, start)
        if idx == -1:
            return False
        before_ok = idx == 0 or not name[idx - 1].isalnum()
        after = idx + len(lot)
        after_ok = after == len(name) or not name[after].isalnum()
        if before_ok and after_ok:
            return True
        start = idx + 1


def classify_lot(lot: str, files: Sequence[DriveFile]) -> Verdict:
    """롯트 하나에 대한 판정.

    ⛔ 넓혀 찾지 않는다. 롯트 문자열이 파일명에 **그대로** 든 것만 받는다 —
       E07Z083 에 E07Z082 를, 416022 에 416006 을 주는 것이 최악의 실패다.
    """
    lot = (lot or "").strip()
    if not lot:
        return Verdict(NONE, (), "롯트가 비어 있습니다")

    # ⛔ 대소문자를 가리지 마라 — Drive 의 contains 는 무시하므로 조회는 파일을
    #    찾아오는데 여기서 도로 버린다. 롯트는 사람이 엑셀에 적은 값이라
    #    소문자로 적었다는 이유로 '없음' 이 나오면 그게 조용한 오답이다
    folded = lot.casefold()
    candidates = [f for f in files if folded in f.name.casefold()]
    if candidates:
        delimited = [f for f in candidates if _is_delimited_match(f.name, lot)]
        if delimited:
            exact = _dedup(delimited)
            if len(exact) > 1:
                return Verdict(MANY, exact, f"{len(exact)}건 — 어느 것인지 확인이 필요합니다")
            if len(lot) <= _SHORT_LOT_LEN:
                return Verdict(CHECK, exact,
                               f"롯트가 {len(lot)}자로 짧아 다른 코드에 걸렸을 수 있습니다")
            return Verdict(FOUND, exact)

        # 앞뒤가 다른 글자와 붙어 있는 매칭뿐이다 — 더 긴 롯트의 일부일 수 있으니
        # 없다고 단정하지 않고 사람에게 넘긴다
        exact = _dedup(candidates)
        return Verdict(
            CHECK, exact,
            "롯트 앞뒤가 다른 글자와 붙어 있습니다 — 더 긴 롯트의 일부일 수 있습니다")

    # 접미가 붙은 롯트("F31C28 D")인데 접미 없는 파일만 있는 경우
    head = lot.split()[0]
    if head != lot:
        head_folded = head.casefold()
        base_hits = _dedup([f for f in files if head_folded in f.name.casefold()])
        if base_hits:
            return Verdict(
                CHECK, base_hits,
                f"접미 없이 '{head}' 로만 된 파일입니다 — 같은 롯트인지 확인하세요")

    # ⛔ note 를 비워 두지 마라 — 화면에 '없음' 세 글자만 남으면 "정말 없다" 와
    #    "토큰이 죽었다"·"드라이브 멤버가 아니다"·"롯트를 잘못 적었다" 가 글자
    #    그대로 똑같이 보인다. 이 기능의 가치는 '없음' 을 믿을 수 있다는 것이다
    return Verdict(NONE, (), _NONE_NOTE)


import concurrent.futures
import logging

from app.core.google_workspace import search_drive

logger = logging.getLogger(__name__)

# 모든 제품에 공통으로 들어가 변별력이 없는 낱말
_BRAND_WORDS = {"skin1004", "madagascar", "centella", "cpnp", "twin", "pack"}
_TERM_SPLIT = re.compile(r"[\s_/,()]+")


@dataclass(frozen=True)
class Result:
    row: Row
    coa: Verdict
    msds: Verdict
    # 제품명으로 찾은 COA. ⛔ 이 열은 **다른 롯트**의 증명서다 — 절대 FOUND 가 아니다
    product_coa: Verdict


def product_terms(description: str) -> list[str]:
    """MSDS 검색에 쓸 제품 고유어. 브랜드어를 빼지 않으면 전 제품이 걸린다."""
    out = []
    for token in _TERM_SPLIT.split(description or ""):
        t = token.strip()
        if len(t) < 2 or t.casefold() in _BRAND_WORDS:
            continue
        out.append(t)
    return out


def _to_files(raw: Iterable[dict]) -> list[DriveFile]:
    return [
        DriveFile(id=r.get("id", ""), name=r.get("name", ""),
                  size=int(r.get("size") or 0), web_link=r.get("webViewLink", ""))
        for r in raw
    ]


def _doc_name_matches(name: str, terms: Sequence[str], marker: str) -> bool:
    """파일명이 표식(msds·coa)과 검색에 쓴 낱말을 전부 담고 있는지 본다.

    ⛔ search_drive 의 fullText 매칭이 열려 있어, 검색어를 본문 어딘가에만
       가진 무관한 문서(인증 서류 리스트, 등록 현황표 등)가 걸린다. 파일명으로
       다시 좁혀 그런 잡음을 걷어낸다 — 없다고 답하는 편이 잡음을 답으로 주는
       것보다 낫다.
    """
    lowered = name.casefold()
    if marker not in lowered:
        return False
    return all(t.casefold() in lowered for t in terms)


def _filter_by_name(files: Iterable[DriveFile], terms: Sequence[str],
                    marker: str) -> list[DriveFile]:
    return [f for f in files if _doc_name_matches(f.name, terms, marker)]


def _one(creds, row: Row, search) -> Result:
    # COA — 롯트 정확 매칭. exact_name 으로 Drive 의 토큰 매칭 잡음을 걷어낸다
    # ⚠️ try 는 네트워크 호출만 감싼다 — classify_lot 은 순수 함수라 여기서
    #    터지면 진짜 버그다. 조회 실패로 위장해 삼키지 않는다
    if row.lot:
        # ⛔ try 안에서 계산하지 마라 — 공백뿐인 롯트면 `split()[0]` 이 IndexError 를
        #    내고 그것이 '조회실패' 로 위장된다. try 는 네트워크 호출만 감싼다
        parts = row.lot.split()
        head = parts[0] if parts else row.lot
        try:
            raw = search(creds, row.lot, max_results=25, exact_name=row.lot)
            if not raw and head != row.lot:
                # exact_name=row.lot 은 접미까지 통째로 요구한다 — 접미 없는
                # 파일도 후보로 올려야 classify_lot 의 확인필요 분기가 실제로
                # 동작한다. 여기서도 exact_name 을 넘겨 본문에만 롯트가 스친
                # 무관 문서(재고 현황표 등)를 걸러낸다
                raw = search(creds, head, max_results=25, exact_name=head)
        except Exception as exc:                      # noqa: BLE001
            # ⚠️ 실패를 삼키지 마라 — 프로덕션은 INFO 를 버린다
            logger.warning("coa_finder_query_failed",
                           extra={"lot": row.lot, "error": str(exc)[:200]})
            coa = Verdict(FAILED, (), _QUERY_FAILED_NOTE + _failure_cause(exc))
        else:
            coa = classify_lot(row.lot, _to_files(raw))
    else:
        coa = Verdict(NONE, (), "롯트가 비어 있습니다")

    # MSDS — 롯트가 없다. 제품 고유어 + 'MSDS' 로 찾는다
    terms = product_terms(row.description)
    if terms:
        query_terms = terms[:4]
        try:
            raw = search(creds, " ".join(query_terms + ["MSDS"]), max_results=25,
                         widen=False)
        except Exception as exc:                      # noqa: BLE001
            logger.warning("msds_finder_query_failed",
                           extra={"sku": row.sku, "error": str(exc)[:200]})
            msds = Verdict(FAILED, (), _QUERY_FAILED_NOTE + _failure_cause(exc))
        else:
            files = _dedup(_filter_by_name(_to_files(raw), query_terms, "msds"))
            if not files:
                msds = Verdict(NONE, (), "")
            elif len(files) == 1:
                msds = Verdict(FOUND, files, "제품명으로 찾았습니다 (롯트 무관)")
            else:
                msds = Verdict(MANY, files, "제품명으로 찾았습니다 (롯트 무관)")
    elif row.description:
        msds = Verdict(NONE, (), "제품명에서 검색에 쓸 낱말을 찾지 못했습니다")
    else:
        msds = Verdict(NONE, (), "제품명이 비어 있습니다")

    return Result(row=row, coa=coa, msds=msds,
                  product_coa=_product_coa(creds, row, search, coa))


def _product_coa(creds, row: Row, search, coa: Verdict) -> Verdict:
    """제품명으로 COA 를 찾는다 — 이 롯트의 것이 **아닌** 증명서를 보여주는 열이다.

    ⛔ 절대 FOUND 를 내지 않는다. 실측(2026-08-31): 요청자의 41행은 롯트로는
       COA 0건인데 34개 제품 중 33개는 COA 가 드라이브에 있다 — 전부 **다른
       생산분**의 것이다. 그러니 아무리 잘 맞아도 "찾았다" 가 될 수 없고,
       가장 세게 말할 수 있는 것이 '확인필요' 다.
    ⚠️ 이 열이 이 기능에서 가장 위험한 자리다. 화면에서 가장 도움이 되어
       보이는 것이 아니라 **가장 조심스러운 것**이어야 한다.
    """
    if coa.status == FOUND:
        # 롯트로 찾았으면 다른 생산분은 필요 없다 — 조회를 아끼되 말은 한다
        return Verdict(NONE, (), "롯트로 COA 를 찾았으므로 제품명으로는 조회하지 않았습니다")

    terms = product_terms(row.description)
    if not terms:
        return Verdict(NONE, (), "제품명이 없어 제품 단위로는 찾을 수 없습니다"
                       if not row.description
                       else "제품명에서 검색에 쓸 낱말을 찾지 못했습니다")

    query_terms = terms[:4]
    try:
        raw = search(creds, " ".join(query_terms + ["COA"]), max_results=25,
                     widen=False)
    except Exception as exc:                          # noqa: BLE001
        logger.warning("product_coa_query_failed",
                       extra={"sku": row.sku, "error": str(exc)[:200]})
        return Verdict(FAILED, (), _QUERY_FAILED_NOTE + _failure_cause(exc))

    files = _dedup(_filter_by_name(_to_files(raw), query_terms, "coa"))
    if not files:
        return Verdict(NONE, (), "이 제품 이름으로도 COA 를 찾지 못했습니다 "
                                 "(파일명 기준 · 파일 본문은 검색하지 않습니다)")
    lot_label = f"({row.lot})" if row.lot else ""
    return Verdict(
        CHECK, files,
        f"제품명으로만 찾은 COA {len(files)}건입니다 — 이 롯트{lot_label}의 것이 "
        "아니라 다른 생산분의 증명서입니다. 쓰기 전에 반드시 열어서 확인하세요")


def find_all(creds, rows: Sequence[Row], search=None, max_workers: int = 8):
    """행 순서를 지키면서 병렬로 조회한다.

    ⛔ `with ThreadPoolExecutor` 로 감싸지 마라 — 블록을 나갈 때 shutdown(wait=True)
       가 걸려 타임아웃이 무의미해진다 (프로젝트 규칙).
    """
    search = search or search_drive
    pool = concurrent.futures.ThreadPoolExecutor(max_workers=max_workers)
    try:
        futures = [pool.submit(_one, creds, row, search) for row in rows]
        for fut in futures:                  # 제출 순서 = 행 순서
            yield fut.result()
    finally:
        pool.shutdown(wait=False)
