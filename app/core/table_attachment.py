# -*- coding: utf-8 -*-
"""올린 엑셀·CSV 를 **붙여넣은 표와 같은 모양(TSV)** 으로 바꾼다.

붐따 #161 (2026-09-03, 정다운):

    "이미지파일만 업데이트가 가능한데, 요거 엑셀 데이터는 못읽는거죠?
     이미지로하니깐 제대로 못 읽는것같아서요!"

맞았다. 채팅은 PNG·JPEG·GIF·WebP 만 받았다. 그래서 표를 **사진으로 찍어** 올렸고,
vision LLM 이 50건을 눈으로 읽었다 — 맞았는지 아무도 확인할 수 없는 값이다.

⛔ **왜 TSV 인가**: 이 앱에는 이미 「사용자가 붙여넣은 표」 경로가 있다
   (`orchestrator._has_pasted_data`). 그 경로는 붙여넣은 표를 조회로 덮어쓰지
   않도록 지키고 있다(2026-08-18 실측 사고). 파일을 **그 사람이 붙여넣었을 때와
   똑같은 문자열**로 바꿔 넣으면, 새 경로를 만들지 않고 그 방어를 그대로 물려받는다.
   ⚠️ 그래서 구분자는 반드시 **탭**이다 — `_INLINE_TSV_ROW` 가 탭을 본다.

⛔ **읽은 것을 숨기지 않는다.** 변환한 표는 대화에 그대로 실려서, 잘못 읽혔으면
   사용자가 **화면에서 바로 본다**. 사진으로 찍어 올렸을 때 없던 성질이다.

⚠️ 첫 시트만 읽는다 — 다른 시트가 있으면 **이름을 함께 알려준다**. 조용히 하나만
   읽으면 나머지가 없는 줄 안다 (COA 찾기가 같은 제약을 지고 있고, 거기서도
   말한다).
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import List, Optional

import structlog

logger = structlog.get_logger(__name__)

MAX_BYTES = 5 * 1024 * 1024      # COA 찾기와 같은 상한
MAX_ROWS = 300                   # 그 이상은 잘라내고 **잘랐다고 말한다**
MAX_COLS = 40
MAX_CHARS = 40_000               # 프롬프트를 지키는 절대 상한

# 본문에 실릴 때 붙는 표식. ⛔ **서버가 블록 전체를 만든다** — 프론트가 같은
#    문자열을 따로 조립하면 사본이 갈리고, 갈리는 순간 아래 판정이 조용히
#    풀린다 (`@@` 목록이 두 벌이라 어긋났던 그 사고와 같은 종류).
ATTACHMENT_MARKER = "[첨부한 표 — "

_XLSX = (".xlsx", ".xlsm")
_LEGACY_XLS = (".xls",)
_TEXTUAL = (".csv", ".tsv", ".txt")


class UnsupportedFile(ValueError):
    """읽을 수 없는 파일. 메시지는 **무엇을 하면 되는지**까지 말한다."""


@dataclass
class ParsedTable:
    rows: List[List[str]]
    sheet: str = ""
    other_sheets: List[str] = field(default_factory=list)
    truncated_rows: int = 0
    truncated_cols: int = 0
    total_rows: int = 0
    total_cols: int = 0

    def to_tsv(self) -> str:
        """탭 구분 텍스트. ⚠️ 칸 안의 탭·줄바꿈은 공백으로 바꾼다 — 안 그러면
        행이 어긋나 표가 조용히 밀린다."""
        out = []
        for row in self.rows:
            out.append("\t".join(_flatten(c) for c in row))
        return "\n".join(out)

    def to_block(self, filename: str) -> str:
        """본문 앞에 붙일 덩어리. 표식 + 무엇을 읽었는지 + 표.

        ⛔ 표식이 있어야 `orchestrator._has_pasted_data` 가 **크기와 무관하게**
           "사용자가 가져온 표" 로 본다. 그 판정에는 200자 문턱이 있어서, 작은
           표(8행 × 3열)는 표식 없이는 걸리지 않고 **조회로 새어** 묻지도 않은
           데이터가 답으로 나간다 (2026-08-18 실측 사고의 작은 판).
        """
        return f"{ATTACHMENT_MARKER}{self.note(filename)}]" + chr(10) + self.to_tsv()

    def note(self, filename: str) -> str:
        """사람이 읽는 한 줄. ⛔ 자른 것·건너뛴 시트는 **반드시 말한다**."""
        bits = [f"`{filename}`"]
        if self.sheet:
            bits.append(f"시트 `{self.sheet}`")
        bits.append(f"{self.total_rows}행 × {self.total_cols}열")
        note = " · ".join(bits)
        if self.truncated_rows:
            note += f" — 앞 {len(self.rows)}행만 실었습니다 ({self.truncated_rows}행 생략)"
        if self.truncated_cols:
            note += f" · 앞 {MAX_COLS}열만 실었습니다 ({self.truncated_cols}열 생략)"
        if self.other_sheets:
            note += (" · 첫 시트만 읽었습니다 (다른 시트: "
                     + ", ".join(f"`{s}`" for s in self.other_sheets) + ")")
        return note


_WS = re.compile(r"[\t\r\n]+")


def _flatten(v) -> str:
    if v is None:
        return ""
    return _WS.sub(" ", str(v)).strip()


def _trim(rows: List[List[str]]) -> ParsedTable:
    """빈 가장자리를 걷어내고 상한을 건다."""
    rows = [[_flatten(c) for c in r] for r in rows]
    while rows and not any(c for c in rows[-1]):
        rows.pop()
    while rows and not any(c for c in rows[0]):
        rows.pop(0)
    total_rows = len(rows)
    total_cols = max((len(r) for r in rows), default=0)

    truncated_cols = max(0, total_cols - MAX_COLS)
    if truncated_cols:
        rows = [r[:MAX_COLS] for r in rows]
    truncated_rows = max(0, total_rows - MAX_ROWS)
    if truncated_rows:
        rows = rows[:MAX_ROWS]

    # 절대 상한 — 넓은 표는 행 수가 적어도 프롬프트를 넘길 수 있다
    while rows and sum(len("\t".join(r)) + 1 for r in rows) > MAX_CHARS:
        rows.pop()
        truncated_rows += 1

    return ParsedTable(rows=rows, truncated_rows=truncated_rows,
                       truncated_cols=truncated_cols,
                       total_rows=total_rows, total_cols=total_cols)


def _parse_xlsx(data: bytes) -> ParsedTable:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:                    # noqa: BLE001
        logger.warning("table_attachment_xlsx_failed", error=str(e)[:160])
        raise UnsupportedFile(
            "엑셀 파일을 열지 못했습니다. 파일이 손상되었거나 암호가 걸려 있을 수 "
            "있습니다 — 표를 복사해 채팅창에 붙여넣어 주세요.") from e
    names = list(wb.sheetnames)
    ws = wb[names[0]]
    # ⚠️ `data_only=True` 는 **엑셀이 계산해 저장해 둔 값**을 준다. 수식만 있고
    #    저장된 값이 없으면 None 이 온다 — 그건 우리가 계산해 줄 수 없다
    raw = [list(r) for r in ws.iter_rows(values_only=True)]
    table = _trim(raw)
    table.sheet = names[0]
    table.other_sheets = names[1:]
    return table


def _decode(data: bytes) -> str:
    """⚠️ 한국어 엑셀이 내보낸 CSV 는 **cp949** 인 경우가 흔하다. utf-8 로만
    읽으면 터지거나 한글이 깨진다 — 이 프로젝트 콘솔이 겪는 것과 같은 함정."""
    for enc in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_text(data: bytes) -> ParsedTable:
    text = _decode(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        # ⚠️ 못 알아보면 쉼표로 본다 — 한 열짜리 목록도 그대로 읽힌다
        delimiter = "\t" if "\t" in sample else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return _trim([row for row in reader])


def looks_like_table(filename: str) -> bool:
    low = (filename or "").lower()
    return low.endswith(_XLSX + _LEGACY_XLS + _TEXTUAL)


def parse(filename: str, data: bytes) -> ParsedTable:
    """파일을 표로 읽는다. 못 읽으면 `UnsupportedFile`.

    ⛔ 예외 메시지는 그대로 사용자에게 간다 — **무엇이 안 되는지가 아니라
       무엇을 하면 되는지** 적는다 (COA 찾기와 같은 규칙).
    """
    low = (filename or "").lower()
    if len(data) > MAX_BYTES:
        raise UnsupportedFile(
            f"파일이 {MAX_BYTES // (1024 * 1024)}MB 상한을 넘습니다 "
            f"(현재 {len(data) / (1024 * 1024):.1f}MB). 필요한 시트·행만 남겨 "
            "다시 올려 주세요.")
    if low.endswith(_LEGACY_XLS):
        # ⛔ openpyxl 은 옛 .xls 를 못 읽는다. "열지 못했습니다" 로 뭉개면
        #    사용자는 파일이 깨진 줄 안다 — 형식 문제라고 말해 준다
        raise UnsupportedFile(
            "옛 형식(.xls)은 읽지 못합니다. 엑셀에서 `.xlsx` 로 저장한 뒤 다시 "
            "올리거나, 표를 복사해 채팅창에 붙여넣어 주세요.")
    if low.endswith(_XLSX):
        table = _parse_xlsx(data)
    elif low.endswith(_TEXTUAL):
        table = _parse_text(data)
    else:
        raise UnsupportedFile(
            "읽을 수 있는 표 파일이 아닙니다 (.xlsx · .csv · .tsv 를 받습니다).")
    if not table.rows:
        raise UnsupportedFile(
            "파일에서 읽을 행을 찾지 못했습니다. 첫 시트가 비어 있는지 "
            "확인해 주세요.")
    return table
